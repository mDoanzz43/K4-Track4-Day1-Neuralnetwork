"""Lưu kết quả JSON và sinh experiments.xlsx từ template.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    exp_id = result["cfg"]["exp_id"]
    out_dir = Path(results_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{exp_id}.json"
    payload = {key: result[key] for key in ("cfg", "history", "summary")}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    directory = Path(results_dir)
    if not directory.exists():
        return []
    results = []
    for path in sorted(directory.glob("*.json")):
        item = json.loads(path.read_text(encoding="utf-8"))
        # results/ còn chứa lock/eval/error-analysis JSON; chỉ lấy JSON của run.
        if all(key in item for key in ("cfg", "history", "summary")):
            results.append(item)
    return sorted(results, key=lambda item: item["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg, summary = result["cfg"], result["summary"]
    row = {
        "exp_id": cfg["exp_id"], "group": cfg.get("group", ""),
        "description": cfg.get("description", ""), "loss": cfg["loss"],
        "optimizer": cfg["optimizer"], "lr": cfg["lr"],
        "weight_decay": cfg.get("weight_decay", 0.0), "batch": cfg["batch"],
        "epochs": cfg["epochs"], "hidden": "-".join(map(str, cfg["hidden"])),
        "dropout": cfg["dropout"], "clip_norm": cfg.get("clip_norm"),
        "precision": cfg["precision"], "init": cfg["init"], "seed": cfg["seed"],
        **summary,
        "eval_acc": None, "eval_macro_f1": None,
        "figure_file": f"figures/{cfg['exp_id']}.png", "notes": notes,
    }
    if eval_scores:
        row["eval_acc"] = eval_scores.get("accuracy", eval_scores.get("acc"))
        row["eval_macro_f1"] = eval_scores.get("macro_f1")
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    formula_columns = {
        "step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"
    }
    workbook = openpyxl.load_workbook(template_path)
    sheet = workbook["Experiments"]
    headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
    for offset, row in enumerate(rows, start=2):
        for key, value in row.items():
            if key in formula_columns or key not in headers:
                continue
            sheet.cell(row=offset, column=headers[key], value=value)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(out_path)
