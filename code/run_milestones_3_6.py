"""Chạy tuần tự Mốc 3–6 trên GPU, hiển thị tiến độ và khóa val trước eval.

Chạy từ repo root:
    venv\Scripts\python.exe code\run_milestones_3_6.py
"""
from __future__ import annotations

import json
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import torch
from openpyxl import load_workbook

from data import prepare_data
from plots import plot_compare, plot_run
from results_table import save_result, to_row, write_xlsx
from train import DEFAULT_CFG, final_eval, run_experiment


HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parent if (HERE.parent / "data").is_dir() else HERE.parents[2]
OUT_DIR = REPO_ROOT / "submission_2A202602839"
RESULTS_DIR = OUT_DIR / "results"
FIGURES_DIR = OUT_DIR / "figures"


def run_one(data: dict, cfg: dict) -> dict:
    """Chạy, lưu JSON/PNG ngay và in một dòng tóm tắt chống mất kết quả."""
    result = run_experiment(cfg, data)
    save_result(result, str(RESULTS_DIR))
    plot_run(result, str(FIGURES_DIR / f"{cfg['exp_id']}.png"))
    s = result["summary"]
    print(
        f"DONE {cfg['exp_id']}: best_epoch={s['best_epoch']} "
        f"val_f1={s['val_macro_f1']:.6f} val_acc={s['val_acc']:.6f} "
        f"sec/epoch={s['time_per_epoch_s']:.2f} diverged={s['diverged']}",
        flush=True,
    )
    return result


def evaluate_predictions(pred_path: Path, out_path: Path) -> dict:
    subprocess.run(
        [sys.executable, "scripts/evaluate.py", "--pred", str(pred_path), "--out", str(out_path)],
        cwd=REPO_ROOT, check=True,
    )
    return json.loads(out_path.read_text(encoding="utf-8"))


def plot_confusion(eval_result: dict, path: Path) -> None:
    cm = np.asarray(eval_result["confusion_matrix"])
    row_sum = cm.sum(axis=1, keepdims=True)
    normalized = np.divide(cm, row_sum, where=row_sum != 0)
    fig, ax = plt.subplots(figsize=(7, 6))
    image = ax.imshow(normalized, cmap="Blues", vmin=0, vmax=1)
    for i in range(7):
        for j in range(7):
            ax.text(j, i, f"{normalized[i, j]:.2f}\n({cm[i, j]})",
                    ha="center", va="center", fontsize=7,
                    color="white" if normalized[i, j] > .55 else "black")
    ax.set(title="Final eval confusion matrix (row-normalized)",
           xlabel="Predicted class", ylabel="True class",
           xticks=range(7), yticks=range(7))
    fig.colorbar(image, ax=ax, fraction=.046, pad=.04)
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)


def fill_workbook(rows: list[dict], baseline_ids: list[str], group_notes: dict[str, str]) -> None:
    target = OUT_DIR / "experiments.xlsx"
    write_xlsx(rows, str(REPO_ROOT / "templates" / "experiment_table_template.xlsx"), str(target))
    workbook = load_workbook(target)
    seeds = workbook["Seeds"]
    for row_idx in range(2, 7):
        seeds.cell(row_idx, 1, baseline_ids[row_idx - 2] if row_idx - 2 < len(baseline_ids) else None)
    summary = workbook["Summary"]
    for row_idx in range(2, summary.max_row + 1):
        group = summary.cell(row_idx, 1).value
        if group in group_notes:
            summary.cell(row_idx, 8, group_notes[group])
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(target)


def main() -> None:
    if not torch.cuda.is_available():
        raise RuntimeError("Cần CUDA để chạy script Mốc 3–6.")
    torch.set_float32_matmul_precision("high")
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    print(f"GPU: {torch.cuda.get_device_name(0)}; torch={torch.__version__}; CUDA={torch.version.cuda}")
    data = prepare_data("cuda", processed_dir=str(REPO_ROOT / "data" / "processed"))

    common = {
        **DEFAULT_CFG, "epochs": 20, "batch": 512, "seed": 1,
        "progress": True, "train_eval_size": 50_000,
    }
    all_results: list[dict] = []

    # Learning-rate search: full-length and same seed for a fair baseline choice.
    lr_results = []
    for lr in (0.03, 0.05, 0.1):
        lr_results.append(run_one(data, {
            **common, "exp_id": f"hparam-lr{str(lr).replace('.', 'p')}", "group": "hparam",
            "description": f"SGD momentum learning-rate search, lr={lr}", "lr": lr,
        }))
    best_lr_result = max(lr_results, key=lambda r: r["summary"]["val_macro_f1"])
    best_lr = float(best_lr_result["cfg"]["lr"])
    print(f"LOCK baseline lr from validation: {best_lr:g}")

    baseline_results = []
    for seed in (1, 2, 3):
        result = run_one(data, {
            **common, "exp_id": f"base-s{seed}", "group": "baseline",
            "description": f"M-base baseline seed {seed}", "lr": best_lr, "seed": seed,
        })
        baseline_results.append(result)
        all_results.append(result)
    baseline = baseline_results[0]
    all_results.extend(lr_results)

    # Loss: change CE -> MSE only.
    all_results.append(run_one(data, {
        **common, "exp_id": "loss-mse", "group": "loss",
        "description": "MSE on one-hot targets versus CE baseline", "lr": best_lr, "loss": "mse",
    }))

    # Optimizers: at least two learning rates per adaptive optimizer.
    optimizer_results = []
    for optimizer_name, lrs in (("adam", (3e-4, 1e-3)), ("adamw", (3e-4, 1e-3))):
        for lr in lrs:
            exp_id = f"opt-{optimizer_name}-lr{lr:g}".replace(".", "p")
            result = run_one(data, {
                **common, "exp_id": exp_id, "group": "optimizer",
                "description": f"{optimizer_name} learning-rate comparison",
                "optimizer": optimizer_name, "lr": lr,
                "weight_decay": 0.01 if optimizer_name == "adamw" else 0.0,
            })
            optimizer_results.append(result)
            all_results.append(result)

    # One-factor experiments for remaining topics.
    hparam_batch = run_one(data, {
        **common, "exp_id": "hparam-batch2048", "group": "hparam",
        "description": "Batch 2048 versus baseline batch 512", "lr": best_lr, "batch": 2048,
    })
    dropout_result = run_one(data, {
        **common, "exp_id": "drop-0p3", "group": "dropout",
        "description": "Dropout q=0.3 versus no dropout", "lr": best_lr, "dropout": 0.3,
    })
    amp_result = run_one(data, {
        **common, "exp_id": "amp-fp16", "group": "amp",
        "description": "CUDA FP16 autocast plus GradScaler", "lr": best_lr, "precision": "fp16",
    })
    init_xavier = run_one(data, {
        **common, "exp_id": "init-xavier", "group": "init",
        "description": "Xavier versus He initialization", "lr": best_lr, "init": "xavier",
    })
    init_zeros = run_one(data, {
        **common, "exp_id": "init-zeros", "group": "init",
        "description": "All-zero weights: symmetry failure check", "lr": best_lr, "init": "zeros",
    })
    all_results.extend((hparam_batch, dropout_result, amp_result, init_xavier, init_zeros))

    baseline_grad = np.asarray(baseline["history"]["grad_norm"])
    clip_c = float(max(0.25, np.median(baseline_grad)))
    high_lr = max(0.5, best_lr * 10)
    clip_results = []
    for clip_norm, suffix in ((None, "noclip"), (clip_c, "clip")):
        result = run_one(data, {
            **common, "exp_id": f"clip-highlr-{suffix}", "group": "clipping",
            "description": f"High lr={high_lr:g}; {suffix}; c={clip_c:.3f}",
            "lr": high_lr, "clip_norm": clip_norm,
        })
        clip_results.append(result)
        all_results.append(result)

    # Select only by validation; exclude deliberately pathological controls.
    final_candidates = [baseline, hparam_batch, dropout_result, amp_result, init_xavier, *optimizer_results]
    selected = max(final_candidates, key=lambda r: r["summary"]["val_macro_f1"])
    locked_cfg = dict(selected["cfg"])
    lock_payload = {
        "selection_rule": "highest val macro-F1 among non-pathological 20-epoch candidates",
        "selected_exp_id": locked_cfg["exp_id"], "cfg": locked_cfg,
        "val_macro_f1": selected["summary"]["val_macro_f1"],
        "val_acc": selected["summary"]["val_acc"],
        "locked_before_eval": True,
    }
    (RESULTS_DIR / "final_config_lock.json").write_text(
        json.dumps(lock_payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("FINAL CONFIG LOCKED BY VAL:", json.dumps(lock_payload, ensure_ascii=False), flush=True)

    # Extra final seeds measure stability but do not alter the locked configuration.
    final_seed_results = []
    for seed in (2, 3):
        cfg = {
            **locked_cfg, "exp_id": f"final-s{seed}", "group": "final",
            "description": f"Locked final configuration, seed {seed}", "seed": seed,
        }
        result = run_one(data, cfg)
        final_seed_results.append(result)
        all_results.append(result)

    groups: dict[str, list[dict]] = defaultdict(list)
    for result in all_results:
        groups[result["cfg"]["group"]].append(result)
    groups["optimizer_with_baseline"] = [baseline, *optimizer_results]
    groups["dropout_with_baseline"] = [baseline, dropout_result]
    groups["amp_with_baseline"] = [baseline, amp_result]
    groups["init_with_baseline"] = [baseline, init_xavier, init_zeros]
    groups["clipping"] = clip_results
    for name, results in groups.items():
        if len(results) >= 2:
            plot_compare(results, "val_macro_f1", str(FIGURES_DIR / f"compare_{name}.png"),
                         title=f"{name}: validation macro-F1")

    # Eval begins only after lock file exists. Evaluate baseline and exactly the locked final model.
    baseline_pred = RESULTS_DIR / "baseline_predictions_eval.csv"
    final_pred = OUT_DIR / "predictions_eval.csv"
    final_eval(baseline["cfg"], baseline, data, str(baseline_pred))
    final_eval(selected["cfg"], selected, data, str(final_pred))
    baseline_eval = evaluate_predictions(baseline_pred, RESULTS_DIR / "eval_baseline.json")
    final_scores = evaluate_predictions(final_pred, OUT_DIR / "eval_result.json")
    baseline_pred.unlink(missing_ok=True)

    cm = np.asarray(final_scores["confusion_matrix"])
    per_class = final_scores["per_class"]
    hardest = min(per_class, key=lambda row: row["f1"])
    hardest_cls = int(hardest["cls"])
    off_diag = cm[hardest_cls].copy()
    off_diag[hardest_cls] = -1
    confused_with = int(off_diag.argmax())
    error_analysis = {
        "hardest_class": hardest_cls, "hardest_class_metrics": hardest,
        "most_confused_with": confused_with,
        "confusion_count": int(cm[hardest_cls, confused_with]),
        "baseline_eval_macro_f1": baseline_eval["macro_f1"],
        "final_eval_macro_f1": final_scores["macro_f1"],
        "eval_improvement": final_scores["macro_f1"] - baseline_eval["macro_f1"],
    }
    (RESULTS_DIR / "error_analysis.json").write_text(
        json.dumps(error_analysis, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    plot_confusion(final_scores, FIGURES_DIR / "confusion_matrix_eval.png")

    eval_by_id = {
        baseline["cfg"]["exp_id"]: baseline_eval,
        selected["cfg"]["exp_id"]: final_scores,
    }
    rows = [to_row(result, eval_by_id.get(result["cfg"]["exp_id"])) for result in all_results]
    group_notes = {
        "baseline": "Ba seed để đo nhiễu; LR được chọn từ validation.",
        "loss": "So CE/MSE bằng accuracy và macro-F1, không so trực tiếp trị số loss.",
        "optimizer": "Adam và AdamW đều được thử hai learning rate.",
        "hparam": "Thử learning rate và batch size; cùng 20 epoch.",
        "dropout": "Đổi duy nhất q từ 0 sang 0,3.",
        "clipping": f"Phản chứng ở lr={high_lr:g}, c={clip_c:.3f} dựa trên grad norm baseline.",
        "amp": "So FP32 baseline với FP16, đo thời gian và peak VRAM.",
        "init": "So He, Xavier và zeros; zeros kiểm tra phá vỡ đối xứng.",
    }
    fill_workbook(rows, [r["cfg"]["exp_id"] for r in baseline_results], group_notes)
    print("EVAL COMPLETE:", json.dumps(error_analysis, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
